"""app/routers/admin.py — Dashboard, settings, search, audit-log, calendar, notifications."""
from __future__ import annotations

import calendar as _calendar
import io
import json
from datetime import date as _date
from typing import Optional

from fastapi import APIRouter, Form, HTTPException, Request
from fastapi.responses import HTMLResponse, RedirectResponse, Response, StreamingResponse
from openpyxl import Workbook

from app.auth import invalidate_all_sessions, require_admin, require_user
from app.db import db, log_action
from app.deps import (
    _build_audit_rows,
    _get_app_settings,
    _is_accounting_user,
    _save_app_settings,
    _validate_iso_date_or_400,
    _validate_time_or_none,
    require_accounting_calendar_access,
    templates,
)

router = APIRouter()


# ── Dashboard ─────────────────────────────────────────────────────────────────

@router.get("/", response_class=HTMLResponse)
def dashboard(request: Request) -> Response:
    import json as _json
    from datetime import timedelta as _td
    today_date = _date.today()

    def _safe_float(v) -> float:
        try:
            return float(str(v or "0").replace(",", ""))
        except (ValueError, TypeError):
            return 0.0

    with db() as conn:
        uploads_total  = conn.execute("SELECT COUNT(*) AS n FROM uploaded_files").fetchone()["n"]
        tx_total       = conn.execute("SELECT COUNT(*) AS n FROM bank_transactions").fetchone()["n"]
        matched_total  = conn.execute("SELECT COUNT(*) AS n FROM bank_transactions WHERE supplier_id IS NOT NULL").fetchone()["n"]
        draft_jes      = conn.execute("SELECT COUNT(*) AS n FROM journal_entries WHERE status='draft'").fetchone()["n"]
        approved_jes   = conn.execute("SELECT COUNT(*) AS n FROM journal_entries WHERE status='approved'").fetchone()["n"]
        je_rejected    = conn.execute("SELECT COUNT(*) AS n FROM journal_entries WHERE status='rejected'").fetchone()["n"]
        suspense_count = conn.execute("SELECT COUNT(*) AS n FROM bank_transactions WHERE classification LIKE '5999%'").fetchone()["n"]

        cash_row = conn.execute("""
            SELECT
              COALESCE(SUM(CAST(REPLACE(debit_amount, ',', '') AS REAL)), 0)  AS total_out,
              COALESCE(SUM(CAST(REPLACE(credit_amount, ',', '') AS REAL)), 0) AS total_in
            FROM bank_transactions
        """).fetchone()
        total_cash_in  = _safe_float(cash_row["total_in"])
        total_cash_out = _safe_float(cash_row["total_out"])
        net_cash       = total_cash_in - total_cash_out

        try:
            pr_row = conn.execute("""
                SELECT COUNT(*) AS n,
                       COALESCE(SUM(CAST(REPLACE(amount,',','') AS REAL)), 0) AS total
                  FROM payment_requests
                 WHERE status IN ('submitted','for_review','for_process')
            """).fetchone()
            pending_requests       = pr_row["n"]
            pending_requests_total = _safe_float(pr_row["total"])
        except Exception:
            pending_requests, pending_requests_total = 0, 0.0

        try:
            inv_row = conn.execute("""
                SELECT COUNT(*) AS n,
                       COALESCE(SUM(CAST(REPLACE(gross_amount,',','') AS REAL)), 0) AS total
                  FROM invoices
                 WHERE status IN ('pending','unmatched','open')
            """).fetchone()
            open_invoices       = inv_row["n"]
            open_invoices_total = _safe_float(inv_row["total"])
        except Exception:
            open_invoices, open_invoices_total = 0, 0.0

        try:
            overdue_count = conn.execute("""
                SELECT COUNT(*) AS n FROM invoices
                 WHERE status IN ('pending','unmatched','open')
                   AND due_date < ?
            """, (today_date.isoformat(),)).fetchone()["n"]
        except Exception:
            overdue_count = 0

        month_start  = today_date.replace(day=1)
        last_m_end   = month_start - _td(days=1)
        last_m_start = last_m_end.replace(day=1)

        tx_this_month = conn.execute("SELECT COUNT(*) AS n FROM bank_transactions WHERE transaction_date >= ?", (month_start.isoformat(),)).fetchone()["n"]
        tx_last_month = conn.execute("SELECT COUNT(*) AS n FROM bank_transactions WHERE transaction_date >= ? AND transaction_date <= ?", (last_m_start.isoformat(), last_m_end.isoformat())).fetchone()["n"]
        je_this_month = conn.execute("SELECT COUNT(*) AS n FROM journal_entries WHERE created_at >= ?", (month_start.isoformat(),)).fetchone()["n"]
        je_last_month = conn.execute("SELECT COUNT(*) AS n FROM journal_entries WHERE created_at >= ? AND created_at <= ?", (last_m_start.isoformat(), last_m_end.isoformat() + " 23:59:59")).fetchone()["n"]

        def _trend(now, prev):
            if prev == 0:
                return None
            pct = round((now - prev) / prev * 100)
            return f"+{pct}%" if pct >= 0 else f"{pct}%"

        stats = {
            "uploads":      uploads_total,
            "transactions": tx_total,
            "matched":      matched_total,
            "draft_jes":    draft_jes,
            "approved_jes": approved_jes,
            "suspense":     suspense_count,
            "tx_trend":     _trend(tx_this_month, tx_last_month),
            "je_trend":     _trend(je_this_month, je_last_month),
            "match_pct":    round(matched_total / tx_total * 100) if tx_total else 0,
        }

        monthly_labels, monthly_debits, monthly_credits, monthly_tx_counts = [], [], [], []
        for i in range(5, -1, -1):
            ref   = today_date.replace(day=1) - _td(days=i * 28)
            ref   = ref.replace(day=1)
            nm    = _calendar.monthrange(ref.year, ref.month)[1]
            m_end = ref.replace(day=nm)
            cf = conn.execute("""
                SELECT
                  COUNT(*) AS n,
                  COALESCE(SUM(CAST(REPLACE(debit_amount,',','') AS REAL)),0)  AS deb,
                  COALESCE(SUM(CAST(REPLACE(credit_amount,',','') AS REAL)),0) AS crd
                FROM bank_transactions
                WHERE transaction_date >= ? AND transaction_date <= ?
            """, (ref.isoformat(), m_end.isoformat())).fetchone()
            monthly_labels.append(ref.strftime("%b %Y"))
            monthly_debits.append(round(_safe_float(cf["deb"]), 2))
            monthly_credits.append(round(_safe_float(cf["crd"]), 2))
            monthly_tx_counts.append(cf["n"])

        recent_transactions = conn.execute("""
            SELECT bt.id, bt.transaction_date, bt.description,
                   bt.counterparty_name, bt.debit_amount, bt.credit_amount,
                   bt.classification, s.name AS supplier_name
              FROM bank_transactions bt
              LEFT JOIN suppliers s ON s.id = bt.supplier_id
             ORDER BY bt.transaction_date DESC, bt.id DESC
             LIMIT 8
        """).fetchall()

        recent_jes = conn.execute(
            "SELECT id, reference AS reference_number, status, created_at FROM journal_entries ORDER BY created_at DESC LIMIT 6"
        ).fetchall()

        recent_uploads = conn.execute(
            "SELECT id, filename, period_covered, parsed_count, uploaded_at FROM uploaded_files ORDER BY uploaded_at DESC LIMIT 5"
        ).fetchall()

        month_last_day = _calendar.monthrange(today_date.year, today_date.month)[1]
        month_end      = today_date.replace(day=month_last_day)
        user           = getattr(request.state, "user", None)
        try:
            if _is_accounting_user(user):
                calendar_events = conn.execute(
                    """SELECT ce.*, u.username, u.full_name FROM calendar_events ce
                       JOIN users u ON u.id = ce.owner_user_id
                       WHERE ce.event_date BETWEEN ? AND ?
                         AND (ce.visibility = 'shared' OR ce.owner_user_id = ?)
                       ORDER BY ce.event_date, COALESCE(ce.start_time,'00:00'), ce.id""",
                    (month_start.isoformat(), month_end.isoformat(), user["id"]),
                ).fetchall()
            else:
                calendar_events = []
        except Exception:
            calendar_events = []

    events_by_date: dict = {}
    for ev in calendar_events:
        events_by_date.setdefault(ev["event_date"], []).append(ev)

    month_days = []
    for week in _calendar.Calendar(firstweekday=6).monthdatescalendar(today_date.year, today_date.month):
        month_days.append([
            {
                "date": day.isoformat(), "day": day.day,
                "in_month": day.month == today_date.month,
                "is_today": day == today_date,
                "events": events_by_date.get(day.isoformat(), []),
            }
            for day in week
        ])

    attention_count = draft_jes + suspense_count + pending_requests + overdue_count

    return templates.TemplateResponse(request, "index.html", {
        "stats":                  stats,
        "recent_uploads":         recent_uploads,
        "recent_jes":             recent_jes,
        "recent_transactions":    recent_transactions,
        "pending_requests":       pending_requests,
        "pending_requests_total": pending_requests_total,
        "open_invoices":          open_invoices,
        "open_invoices_total":    open_invoices_total,
        "overdue_count":          overdue_count,
        "total_cash_in":          total_cash_in,
        "total_cash_out":         total_cash_out,
        "net_cash":               net_cash,
        "attention_count":        attention_count,
        "monthly_labels":         _json.dumps(monthly_labels),
        "monthly_debits":         _json.dumps(monthly_debits),
        "monthly_credits":        _json.dumps(monthly_credits),
        "monthly_values":         _json.dumps(monthly_tx_counts),
        "je_draft":               draft_jes,
        "je_approved":            approved_jes,
        "je_rejected":            je_rejected,
        "calendar_month_name":    today_date.strftime("%B %Y"),
        "calendar_weeks":         month_days,
        "calendar_events":        calendar_events,
    })


# ── Settings ──────────────────────────────────────────────────────────────────

@router.get("/settings", response_class=HTMLResponse)
def settings_page(request: Request, message: str = "", error: str = "") -> Response:
    require_admin(request)
    with db() as conn:
        companies = conn.execute(
            "SELECT * FROM companies ORDER BY is_active DESC, name"
        ).fetchall()
        departments = conn.execute(
            """SELECT d.*, c.name AS company_name
               FROM departments d
               LEFT JOIN companies c ON c.id = d.company_id
               ORDER BY d.is_active DESC, c.name, d.name"""
        ).fetchall()
    return templates.TemplateResponse(request, "settings.html", {
        "companies": companies,
        "departments": departments,
        "app_settings": _get_app_settings(),
        "message": message or None,
        "error": error or None,
    })


@router.post("/settings/companies/create")
def settings_create_company(request: Request, name: str = Form(...)) -> Response:
    require_admin(request)
    name = name.strip()
    if not name:
        return RedirectResponse("/settings?error=Company+name+is+required", status_code=303)
    try:
        with db() as conn:
            cur = conn.execute("INSERT INTO companies (name) VALUES (?)", (name,))
            log_action(conn, "create", "company", cur.lastrowid, {"name": name})
    except Exception:
        return RedirectResponse(f"/settings?error=Company+%22{name}%22+already+exists", status_code=303)
    return RedirectResponse("/settings?message=Company+created", status_code=303)


@router.post("/settings/companies/{company_id}/update")
def settings_update_company(
    request: Request,
    company_id: int,
    name: str = Form(...),
    is_active: str = Form("1"),
) -> Response:
    require_admin(request)
    try:
        with db() as conn:
            conn.execute(
                "UPDATE companies SET name = ?, is_active = ? WHERE id = ?",
                (name.strip(), 1 if is_active == "1" else 0, company_id),
            )
            log_action(conn, "update", "company", company_id, {"name": name, "is_active": is_active})
    except Exception:
        return RedirectResponse("/settings?error=Company+name+already+in+use", status_code=303)
    return RedirectResponse("/settings?message=Company+updated", status_code=303)


@router.post("/settings/departments/create")
def settings_create_department(
    request: Request,
    name: str = Form(...),
    company_id: int = Form(0),
) -> Response:
    require_admin(request)
    name = name.strip()
    if not name:
        return RedirectResponse("/settings?error=Department+name+is+required", status_code=303)
    try:
        with db() as conn:
            cur = conn.execute(
                "INSERT INTO departments (name, company_id) VALUES (?, ?)",
                (name, company_id or None),
            )
            log_action(conn, "create", "department", cur.lastrowid,
                       {"name": name, "company_id": company_id})
    except Exception:
        return RedirectResponse(
            "/settings?error=Department+already+exists+for+that+company", status_code=303
        )
    return RedirectResponse("/settings?message=Department+created", status_code=303)


@router.post("/settings/departments/{dept_id}/update")
def settings_update_department(
    request: Request,
    dept_id: int,
    name: str = Form(...),
    company_id: int = Form(0),
    is_active: str = Form("1"),
) -> Response:
    require_admin(request)
    with db() as conn:
        conn.execute(
            "UPDATE departments SET name = ?, company_id = ?, is_active = ? WHERE id = ?",
            (name.strip(), company_id or None, 1 if is_active == "1" else 0, dept_id),
        )
        log_action(conn, "update", "department", dept_id,
                   {"name": name, "is_active": is_active})
    return RedirectResponse("/settings?message=Department+updated", status_code=303)


@router.post("/settings/rotate-sessions", response_class=HTMLResponse)
def settings_rotate_sessions(request: Request) -> Response:
    """Global emergency: invalidate every active user session immediately.

    Bumps session_version for all active users in one UPDATE.  Every existing
    session cookie — including the admin's own — becomes invalid.  The admin
    will be redirected to /login after their next request.

    Use this when you suspect session tokens have been compromised.
    """
    admin = require_admin(request)
    affected = invalidate_all_sessions(actor=admin["username"])
    # The admin's own session is now invalid.  Redirect to login with a
    # message that explains what happened so they don't think it's a bug.
    return RedirectResponse(
        f"/login?next=/settings&message=All+sessions+rotated.+{affected}+user"
        f"{'s' if affected != 1 else ''}+signed+out.",
        status_code=303,
    )


@router.post("/settings/smtp")
def settings_save_smtp(
    request: Request,
    smtp_host: str = Form(""),
    smtp_port: str = Form("587"),
    smtp_user: str = Form(""),
    smtp_password: str = Form(""),
    smtp_from: str = Form(""),
    smtp_tls: str = Form("0"),
) -> Response:
    require_admin(request)
    current = _get_app_settings()
    settings = {
        "smtp_host": smtp_host.strip(),
        "smtp_port": smtp_port.strip() or "587",
        "smtp_user": smtp_user.strip(),
        "smtp_from": smtp_from.strip(),
        "smtp_tls": "1" if smtp_tls == "1" else "0",
    }
    if smtp_password.strip():
        settings["smtp_password"] = smtp_password.strip()
    elif "smtp_password" in current:
        settings["smtp_password"] = current.get("smtp_password", "")
    _save_app_settings(settings)
    return RedirectResponse("/settings?message=Email+settings+saved", status_code=303)


# ── Global search ─────────────────────────────────────────────────────────────

@router.get("/search", response_class=HTMLResponse)
def search_page(request: Request, q: str = "") -> Response:
    require_user(request)
    q = q.strip()
    results: dict = {}
    if len(q) >= 2:
        like = f"%{q}%"
        with db() as conn:
            suppliers = conn.execute(
                "SELECT id, name, tin, is_active FROM suppliers WHERE name LIKE ? OR tin LIKE ? ORDER BY is_active DESC, name LIMIT 15",
                (like, like),
            ).fetchall()
            transactions = conn.execute(
                "SELECT id, description, transaction_date, debit_amount, credit_amount FROM bank_transactions WHERE description LIKE ? OR reference_number LIKE ? ORDER BY transaction_date DESC LIMIT 20",
                (like, like),
            ).fetchall()
            accounts = conn.execute(
                "SELECT id, code, name, type, is_active FROM chart_of_accounts WHERE code LIKE ? OR name LIKE ? ORDER BY is_active DESC, code LIMIT 15",
                (like, like),
            ).fetchall()
            journal_entries = conn.execute(
                "SELECT je.id, je.reference, je.description, je.entry_date, je.status FROM journal_entries je WHERE je.reference LIKE ? OR je.description LIKE ? ORDER BY je.entry_date DESC LIMIT 15",
                (like, like),
            ).fetchall()
        results = {
            "suppliers": suppliers,
            "transactions": transactions,
            "accounts": accounts,
            "journal_entries": journal_entries,
        }
    return templates.TemplateResponse(request, "search.html", {
        "q": q,
        "results": results,
        "total": sum(len(v) for v in results.values()),
    })


# ── Notifications ─────────────────────────────────────────────────────────────

@router.get("/notifications", response_class=HTMLResponse)
def notifications_page(request: Request) -> Response:
    from app.deps import _current_user
    user = _current_user(request)
    with db() as conn:
        rows = conn.execute(
            "SELECT * FROM notifications WHERE user_id = ? ORDER BY created_at DESC LIMIT 200",
            (user["id"],),
        ).fetchall()
        conn.execute("UPDATE notifications SET is_read = 1 WHERE user_id = ?", (user["id"],))
    return templates.TemplateResponse(request, "notifications.html", {"notifications": rows})


# ── Calendar ──────────────────────────────────────────────────────────────────

@router.get("/calendar", response_class=HTMLResponse)
def calendar_page(request: Request, view: str = "upcoming", message: str = "", edit: Optional[int] = None) -> Response:
    user = require_accounting_calendar_access(request)
    today = _date.today().isoformat()
    # Build visibility filter using only literal SQL fragments; user["id"] is an
    # integer from the authenticated session (not raw user input) and travels
    # through ? placeholders, never into the SQL string itself.
    if view == "private":
        visibility_sql = "visibility = 'private' AND owner_user_id = ?"
        params: list = [user["id"]]
    elif view == "shared":
        visibility_sql = "visibility = 'shared'"
        params = []
    elif view == "all":
        visibility_sql = "(visibility = 'shared' OR owner_user_id = ?)"
        params = [user["id"]]
    else:
        visibility_sql = "(visibility = 'shared' OR owner_user_id = ?) AND event_date >= ?"
        params = [user["id"], today]
        view = "upcoming"
    with db() as conn:
        rows = conn.execute(
            "SELECT ce.*, u.username, u.full_name"
            " FROM calendar_events ce"
            " JOIN users u ON u.id = ce.owner_user_id"
            " WHERE " + visibility_sql +
            " ORDER BY ce.event_date ASC, COALESCE(ce.start_time, '00:00') ASC, ce.id ASC",
            tuple(params),
        ).fetchall()
    return templates.TemplateResponse(request, "calendar.html", {
        "rows": rows,
        "view": view,
        "today": today,
        "message": message or None,
        "visibility_options": ["shared", "private"],
        "edit_id": edit,
    })


@router.post("/calendar/create")
def calendar_create(
    request: Request,
    title: str = Form(...),
    event_date: str = Form(...),
    start_time: str = Form(""),
    end_time: str = Form(""),
    description: str = Form(""),
    location: str = Form(""),
    visibility: str = Form("shared"),
    reminder_minutes: str = Form(""),
) -> Response:
    user = require_accounting_calendar_access(request)
    title = title.strip()
    if not title:
        raise HTTPException(status_code=400, detail="Event title is required")
    event_date = _validate_iso_date_or_400(event_date, "event date")
    start_time_v = _validate_time_or_none(start_time, "start time")
    end_time_v = _validate_time_or_none(end_time, "end time")
    if visibility not in {"shared", "private"}:
        raise HTTPException(status_code=400, detail="Visibility must be shared or private")
    reminder = None
    if reminder_minutes.strip():
        try:
            reminder = int(reminder_minutes)
        except ValueError as exc:
            raise HTTPException(status_code=400, detail="Reminder must be minutes before the event") from exc
        if reminder < 0:
            raise HTTPException(status_code=400, detail="Reminder must be zero or more minutes")
    with db() as conn:
        cur = conn.execute(
            """INSERT INTO calendar_events
               (title, event_date, start_time, end_time, description, location, visibility, owner_user_id, reminder_minutes)
               VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)""",
            (title, event_date, start_time_v, end_time_v, description.strip() or None,
             location.strip() or None, visibility, user["id"], reminder),
        )
        event_id = cur.lastrowid
        log_action(conn, "create", "calendar_event", event_id,
                   {"title": title, "event_date": event_date, "visibility": visibility},
                   user_id=user["username"])
    return RedirectResponse("/calendar?message=Event%20created", status_code=303)


@router.post("/calendar/{event_id}/update")
def calendar_update(
    request: Request,
    event_id: int,
    title: str = Form(...),
    event_date: str = Form(...),
    start_time: str = Form(""),
    end_time: str = Form(""),
    description: str = Form(""),
    location: str = Form(""),
    visibility: str = Form("shared"),
    reminder_minutes: str = Form(""),
    is_done: str = Form("0"),
) -> Response:
    user = require_accounting_calendar_access(request)
    title = title.strip()
    if not title:
        raise HTTPException(status_code=400, detail="Event title is required")
    event_date = _validate_iso_date_or_400(event_date, "event date")
    start_time_v = _validate_time_or_none(start_time, "start time")
    end_time_v = _validate_time_or_none(end_time, "end time")
    if visibility not in {"shared", "private"}:
        raise HTTPException(status_code=400, detail="Visibility must be shared or private")
    reminder = None
    if reminder_minutes.strip():
        try:
            reminder = int(reminder_minutes)
        except ValueError as exc:
            raise HTTPException(status_code=400, detail="Reminder must be minutes before the event") from exc
        if reminder < 0:
            raise HTTPException(status_code=400, detail="Reminder must be zero or more minutes")
    with db() as conn:
        ev = conn.execute("SELECT * FROM calendar_events WHERE id = ?", (event_id,)).fetchone()
        if ev is None:
            raise HTTPException(status_code=404, detail="Calendar event not found")
        if ev["visibility"] == "private" and ev["owner_user_id"] != user["id"] and user["role"] != "admin":
            raise HTTPException(status_code=403, detail="Private event access denied")
        conn.execute(
            """UPDATE calendar_events
               SET title = ?, event_date = ?, start_time = ?, end_time = ?, description = ?, location = ?,
                   visibility = ?, reminder_minutes = ?, is_done = ?, updated_at = datetime('now')
               WHERE id = ?""",
            (title, event_date, start_time_v, end_time_v, description.strip() or None,
             location.strip() or None, visibility, reminder, 1 if is_done == "1" else 0, event_id),
        )
        log_action(conn, "update", "calendar_event", event_id,
                   {"title": title, "event_date": event_date, "visibility": visibility},
                   user_id=user["username"])
    return RedirectResponse("/calendar?message=Event%20updated", status_code=303)


@router.post("/calendar/{event_id}/delete")
def calendar_delete(request: Request, event_id: int) -> Response:
    user = require_accounting_calendar_access(request)
    with db() as conn:
        ev = conn.execute("SELECT * FROM calendar_events WHERE id = ?", (event_id,)).fetchone()
        if ev is None:
            raise HTTPException(status_code=404, detail="Calendar event not found")
        if ev["visibility"] == "private" and ev["owner_user_id"] != user["id"] and user["role"] != "admin":
            raise HTTPException(status_code=403, detail="Private event access denied")
        conn.execute("DELETE FROM calendar_events WHERE id = ?", (event_id,))
        log_action(conn, "delete", "calendar_event", event_id,
                   {"title": ev["title"], "event_date": ev["event_date"]},
                   user_id=user["username"])
    return RedirectResponse("/calendar?message=Event%20deleted", status_code=303)


# ── Audit Log ─────────────────────────────────────────────────────────────────

@router.get("/audit-log", response_class=HTMLResponse)
def audit_log(
    request: Request,
    action: str = "",
    entity_type: str = "",
    user_id: str = "",
    date_from: str = "",
    date_to: str = "",
    page: int = 1,
) -> Response:
    require_admin(request)
    page = max(1, page)
    per_page = 100
    offset = (page - 1) * per_page

    # Build the WHERE clause from a fixed "1=1" base so that no f-string ever
    # receives user-supplied text.  Every filter value travels through a ?
    # placeholder in `params`; only literal SQL fragments are appended to
    # `where_sql`.
    where_sql = "WHERE 1=1"
    params: list = []
    if action:
        where_sql += " AND action = ?"
        params.append(action)
    if entity_type:
        where_sql += " AND entity_type = ?"
        params.append(entity_type)
    if user_id:
        where_sql += " AND lower(user_id) LIKE lower(?)"
        params.append(f"%{user_id}%")
    if date_from:
        where_sql += " AND timestamp >= ?"
        params.append(date_from)
    if date_to:
        where_sql += " AND timestamp <= ?"
        params.append(date_to + " 23:59:59")

    with db() as conn:
        total = conn.execute(
            "SELECT COUNT(*) AS n FROM audit_logs " + where_sql, params
        ).fetchone()["n"]
        raw_rows = conn.execute(
            "SELECT * FROM audit_logs " + where_sql + " ORDER BY id DESC LIMIT ? OFFSET ?",
            params + [per_page, offset],
        ).fetchall()
        actions = [r[0] for r in conn.execute("SELECT DISTINCT action FROM audit_logs ORDER BY action").fetchall()]
        entity_types = [r[0] for r in conn.execute("SELECT DISTINCT entity_type FROM audit_logs ORDER BY entity_type").fetchall()]
        users = [r[0] for r in conn.execute("SELECT DISTINCT user_id FROM audit_logs WHERE user_id IS NOT NULL ORDER BY user_id").fetchall()]

    rows_with_details = _build_audit_rows(raw_rows)
    total_pages = max(1, (total + per_page - 1) // per_page)

    return templates.TemplateResponse(request, "audit_log.html", {
        "rows": rows_with_details,
        "total": total,
        "page": page,
        "total_pages": total_pages,
        "per_page": per_page,
        "filter_action": action,
        "filter_entity_type": entity_type,
        "filter_user_id": user_id,
        "filter_date_from": date_from,
        "filter_date_to": date_to,
        "actions": actions,
        "entity_types": entity_types,
        "users": users,
    })


@router.get("/audit-log/export")
def audit_log_export(
    request: Request,
    action: str = "",
    entity_type: str = "",
    user_id: str = "",
    date_from: str = "",
    date_to: str = "",
) -> Response:
    require_admin(request)
    # Same safe pattern as audit_log(): only literal SQL fragments are appended
    # to where_sql; all user values go through ? placeholders in params.
    where_sql = "WHERE 1=1"
    params: list = []
    if action:
        where_sql += " AND action = ?"
        params.append(action)
    if entity_type:
        where_sql += " AND entity_type = ?"
        params.append(entity_type)
    if user_id:
        where_sql += " AND lower(user_id) LIKE lower(?)"
        params.append(f"%{user_id}%")
    if date_from:
        where_sql += " AND timestamp >= ?"
        params.append(date_from)
    if date_to:
        where_sql += " AND timestamp <= ?"
        params.append(date_to + " 23:59:59")

    with db() as conn:
        raw_rows = conn.execute(
            "SELECT * FROM audit_logs " + where_sql + " ORDER BY id DESC",
            params,
        ).fetchall()

    from app.services.excel_styles import (
        add_corp_header,
        auto_col_width,
        finalize_workbook,
        style_data_rows,
        write_column_headers,
    )
    rows = _build_audit_rows(raw_rows)
    wb = Workbook()
    ws = wb.active
    ws.title = "Audit Log"
    headers = ["ID", "Timestamp", "Action", "Entity Type", "Entity ID", "User", "Details"]
    period = f"{date_from or '—'} to {date_to or '—'}"
    num_cols = len(headers)
    data_row = add_corp_header(ws, "Audit Log", period, num_cols)
    write_column_headers(ws, data_row, headers)
    data_row += 1

    for r in rows:
        details_str = json.dumps(r["details"]) if isinstance(r["details"], dict) else str(r["details"] or "")
        ws.append([r["id"], r["timestamp"], r["action"], r["entity_type"],
                   r["entity_id"], r["user_id"], details_str])

    last_data = ws.max_row
    style_data_rows(ws, data_row, last_data, num_cols)
    auto_col_width(ws)
    ws.column_dimensions["A"].width = 8
    ws.column_dimensions["B"].width = 20
    ws.column_dimensions["G"].width = 50

    fname = f"audit_log_{_date.today().isoformat()}.xlsx"
    return StreamingResponse(
        io.BytesIO(finalize_workbook(wb)),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{fname}"'},
    )
